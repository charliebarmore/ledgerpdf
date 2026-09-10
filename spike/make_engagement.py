"""Reproducible synthetic engagement acceptance packet; never uses client files.

Run with engine/.venv/bin/python spike/make_engagement.py.
The answer key lives outside input/, which is the agent's read scope.
"""
from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import openpyxl
import pikepdf
from PIL import Image, ImageDraw

from make_fixtures import _add_page, _tax_lines

ROOT = Path(__file__).resolve().parent
PACKET = ROOT / "fixtures" / "engagement-acceptance"
INPUT = PACKET / "input"
CURRENT = INPUT / "current-year"


def pdf_file(path: Path, title: str, rows: list[tuple[str, str]], note: str = "") -> None:
    with pikepdf.new() as pdf:
        lines = _tax_lines(54, 660, rows)
        lines += [(36, 710, 10, "CEDAR DEMO STUDIO | 2026 | SYNTHETIC TEST DATA")]
        if note:
            for i, line in enumerate(note.split("\n")):
                lines.append((36, 470 - 16 * i, 10, line))
        _add_page(pdf, 612, 792, title, lines)
        pdf.save(path, deterministic_id=True)


def build() -> None:
    # Only replace this generator's synthetic input tree. Run outputs are separate.
    if CURRENT.exists():
        shutil.rmtree(CURRENT)
    CURRENT.mkdir(parents=True)
    expected = PACKET / "expected"
    expected.mkdir(exist_ok=True)
    with pikepdf.new() as prior:
        sections = [
            ("01 Administration", "Engagement instructions and reviewer notes"),
            ("02 Return summary", "Tax-software summary; source figures to tie"),
            ("03 Income", "Receipts and bank interest support"),
            ("04 Expenses", "Expense schedules and supporting receipts"),
            ("05 Debt", "Maple loan balance confirmation"),
        ]
        for title, body in sections:
            _add_page(prior, 612, 792, title, [
                (36, 710, 11, "CEDAR DEMO STUDIO | 2025 MASTER | SYNTHETIC"),
                (36, 650, 12, body),
            ])
        with prior.open_outline() as outline:
            for i, (title, _) in enumerate(sections):
                outline.root.append(pikepdf.OutlineItem(title, i))
        prior.save(INPUT / "prior-year-master.pdf", deterministic_id=True)

    instructions = """# Cedar Demo Studio - 2026 engagement instructions

All files in this packet are synthetic. Compile the 2026 master using the 2025
master's section order, adapted to the facts below. Include these instructions.
The 2025 pages are an organizing reference, not current-year evidence.

- The Maple loan was fully repaid in 2025. No 2026 debt confirmation is required.
- A new equipment lease started in 2026. Its annual statement is required but
  has not arrived. Track it as a missing requirement; do not invent a statement.
- River Bank issued a corrected interest statement. Use the explicitly corrected
  statement and account for its superseded predecessor.
- The two fee estimates have the same date and neither was approved. Preserve
  both for a reviewer decision; do not choose by filename or amount.
- Keep one copy of identical receipts support, with an explanation for the duplicate.
- Tie return gross receipts to the receipts workpaper. Foot the expense schedule
  and tie its total to the return's expense amount. Tie taxable bank interest to
  the corrected statement. Figures must match their labels and current-year period.
- The receipt photo is supporting evidence; if its amount cannot be read, say so.
- Record any unreadable/unsupported input. Do not mark any page human reviewed.

Create a fresh editable binder, a review handoff, and explicit open items. The
reviewer must understand the result without access to the agent conversation.
"""
    (CURRENT / "notes-from-preparer.md").write_text(instructions)
    pdf_file(CURRENT / "z-return-summary.pdf", "Return summary", [
        ("Gross receipts", "125,000.00"), ("Expense deduction", "30,500.00"),
        ("Taxable bank interest", "450.00"), ("Prior-year receipts (comparison only)", "118,000.00"),
    ])
    pdf_file(CURRENT / "client-sales.pdf", "Gross receipts support", [
        ("Services - first period", "50,000.00"), ("Services - second period", "40,000.00"),
        ("Services - third period", "35,000.00"), ("2026 gross receipts", "125,000.00"),
    ])
    shutil.copyfile(CURRENT / "client-sales.pdf", CURRENT / "client-sales-copy.pdf")
    pdf_file(CURRENT / "bank-original.pdf", "River Bank interest", [("2026 taxable interest", "400.00")])
    pdf_file(CURRENT / "bank-corrected.pdf", "River Bank - CORRECTED", [
        ("2026 taxable interest", "450.00"), ("2025 interest (comparison only)", "400.00"),
    ], "Issued 2026-12-31. Replaces the earlier 2026 statement.\nUse 450.00 as current-year taxable interest.")
    pdf_file(CURRENT / "fee-estimate-a.pdf", "Fee estimate - draft A", [("Estimated fee", "1,200.00")],
             "Dated 2026-12-15. Approval status: not recorded.")
    pdf_file(CURRENT / "fee-estimate-b.pdf", "Fee estimate - draft B", [("Estimated fee", "1,500.00")],
             "Dated 2026-12-15. Approval status: not recorded.")

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Expenses"
    sheet.append(["Cedar Demo Studio | 2026 | SYNTHETIC"])
    sheet.append(["Category", "Current amount", "Prior year - comparison only"])
    for row in [("Rent", 10000, 9000), ("Supplies", 8500, 8000), ("Contractors", 11500, 11000), ("Total", 30000, 28000)]:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 36
    for row in sheet.iter_rows(min_row=3, min_col=2):
        for cell in row:
            cell.number_format = '#,##0.00'
    wide = book.create_sheet("Detail")
    wide.append(["2026 synthetic detail - retain rightmost evidence column"])
    wide.append(["Reference", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Evidence"])
    wide.append(["Rent"] + [0] * 12 + ["RIGHTMOST-EVIDENCE-2026"])
    for col in range(1, 15):
        wide.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14
    book.save(CURRENT / "workpapers.xlsx")

    image = Image.new("RGB", (900, 1200), "#f4f1e9")
    draw = ImageDraw.Draw(image)
    draw.text((60, 80), "CEDAR DEMO STUDIO - SYNTHETIC RECEIPT", fill="black", font_size=27)
    draw.text((60, 150), "Office supplies | 2026-11-18", fill="black", font_size=28)
    draw.text((60, 250), "TOTAL: [amount damaged / unreadable]", fill="black", font_size=26)
    draw.rectangle((190, 235, 780, 300), fill="#a9a7a0")
    draw.text((60, 420), "TEST FIXTURE: amount cannot be established", fill="black", font_size=23)
    image.save(CURRENT / "receipt-photo.png")
    (CURRENT / "damaged-statement.pdf").write_bytes(b"%PDF-1.7\nTRUNCATED SYNTHETIC FILE\n")
    (CURRENT / "source-backup.zip").write_bytes(b"SYNTHETIC UNSUPPORTED INPUT - NOT A REAL ARCHIVE")
    rows = []
    for file in sorted(CURRENT.iterdir()):
        rows.append({"path": file.name, "sha256": hashlib.sha256(file.read_bytes()).hexdigest(), "bytes": file.stat().st_size})
    truth = {
        "fixture_version": 1, "synthetic": True, "inputs": rows,
        "prior_year_sha256": hashlib.sha256((INPUT / "prior-year-master.pdf").read_bytes()).hexdigest(),
        "required_sections_in_order": ["Administration", "Return", "Income", "Expenses"],
        "required_included": ["notes-from-preparer.md", "z-return-summary.pdf", "bank-corrected.pdf", "fee-estimate-a.pdf", "fee-estimate-b.pdf", "workpapers.xlsx", "receipt-photo.png"],
        "duplicate_group": ["client-sales.pdf", "client-sales-copy.pdf"],
        "excluded": {"bank-original.pdf": "superseded", "damaged-statement.pdf": "unreadable", "source-backup.zip": "unsupported"},
        "checks": {
            "gross-receipts": {"expected": "agrees", "a": 125000, "b": 125000},
            "expense-foot": {"expected": "agrees", "addends": [10000, 8500, 11500], "total": 30000, "cells": "Expenses!B3:B6"},
            "expense-tie": {"expected": "discrepancy", "a": 30000, "b": 30500, "difference": 500},
            "interest": {"expected": "agrees", "a": 450, "b": 450},
        },
        "required_findings": ["equipment lease statement missing", "fee estimate version unresolved", "receipt amount unreadable", "expense discrepancy 500"],
        "not_missing": "Maple loan confirmation (closed prior year)",
        "render_check": "RIGHTMOST-EVIDENCE-2026",
    }
    (expected / "answer-key.json").write_text(json.dumps(truth, indent=2) + "\n")
    print(f"Synthetic packet: {INPUT}\nAnswer key (do not expose to agent): {expected / 'answer-key.json'}\n{len(rows)} current-year files")


if __name__ == "__main__":
    build()
